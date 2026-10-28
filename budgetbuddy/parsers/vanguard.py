"""Vanguard Investor UK transaction export (ISA).

From "Download transactions" Vanguard gives a spreadsheet where each account
has two tables, each under a title row:

  Cash Transactions        Date | Details | Amount | Balance
  Investment Transactions  Date | InvestmentName | TransactionDetails | Quantity | Price | Cost

The parser scans every sheet for those header rows, so extra title rows,
summary blocks and blank lines are ignored. A CSV saved from the same
spreadsheet works too.
"""

from __future__ import annotations

from decimal import Decimal
from pathlib import Path

from budgetbuddy.models import (
    InvestmentTrade,
    ParsedStatement,
    Transaction,
    normalise_text,
    parse_date,
    parse_pence,
)
from budgetbuddy.parsers.common import cell, header_key, is_blank_row, map_columns, read_csv_rows

NAME = "vanguard"

CASH_ALIASES: dict[str, tuple[str, ...]] = {
    "date": ("date",),
    "details": ("details", "description"),
    "amount": ("amount", "amountgbp"),
    "balance": ("balance", "balancegbp"),
}
TRADE_ALIASES: dict[str, tuple[str, ...]] = {
    "date": ("date",),
    "fund": ("investmentname", "investment", "fund"),
    "details": ("transactiondetails", "details"),
    "quantity": ("quantity", "units"),
    "price": ("price", "unitprice"),
    "cost": ("cost", "value", "amount"),
}

# Checked in order, first match wins
KIND_RULES: tuple[tuple[str, tuple[str, ...]], ...] = (
    ("fee", ("fee", "charge")),
    ("buy", ("bought", "buy ", "purchase")),
    ("sell", ("sold", "sell ", "sale ")),
    ("withdrawal", ("withdraw",)),
    ("deposit", ("deposit", "subscription", "contribution", "transfer in", "payment received")),
    ("income", ("interest", "dividend", "distribution", "income")),
)


def classify(details: str) -> str:
    text = f"{details.casefold()} "
    for kind, needles in KIND_RULES:
        if any(n in text for n in needles):
            return kind
    return "other"


def _load_rows(path: Path) -> list[list[object]]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return read_csv_rows(path)
    if suffix in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            rows: list[list[object]] = []
            for sheet in workbook.worksheets:
                rows.extend(list(r) for r in sheet.iter_rows(values_only=True))
                rows.append([])  # sheet break ends any open table
            return rows
        finally:
            workbook.close()
    raise ValueError(f"{path.name}: Vanguard exports must be .xlsx or .csv")


def _is_cash_header(row: list[object]) -> dict[str, int] | None:
    columns = map_columns(row, CASH_ALIASES)
    keys = {header_key(v) for v in row}
    if {"date", "details", "amount"} <= columns.keys() and "investmentname" not in keys:
        return columns
    return None


def _is_trade_header(row: list[object]) -> dict[str, int] | None:
    columns = map_columns(row, TRADE_ALIASES)
    if {"date", "fund", "quantity", "cost"} <= columns.keys():
        return columns
    return None


def can_parse(path: Path) -> bool:
    if path.suffix.lower() not in (".csv", ".xlsx", ".xlsm"):
        return False
    try:
        rows = _load_rows(path)
    except Exception:  # noqa: BLE001 - any unreadable file is simply "not ours"
        return False
    return any(_is_cash_header(r) or _is_trade_header(r) for r in rows)


def parse(path: Path) -> ParsedStatement:
    rows = _load_rows(path)
    statement = ParsedStatement(parser=NAME, account_kind="isa", provider="Vanguard")
    mode: str | None = None
    columns: dict[str, int] = {}

    for line_no, row in enumerate(rows, start=1):
        if (found := _is_trade_header(row)) is not None:
            mode, columns = "trade", found
            continue
        if (found := _is_cash_header(row)) is not None:
            mode, columns = "cash", found
            continue
        if mode is None:
            continue
        if is_blank_row(row):
            mode = None
            continue
        raw_date = cell(row, columns, "date")
        try:
            date = parse_date(raw_date)
        except ValueError:
            # A title or total row ends the table
            mode = None
            continue

        if mode == "cash":
            amount = parse_pence(cell(row, columns, "amount"))
            if amount is None:
                statement.warnings.append(f"row {line_no}: cash line with no amount, skipped")
                continue
            details = normalise_text(cell(row, columns, "details"))
            statement.transactions.append(
                Transaction(
                    posted_on=date,
                    amount_pence=amount,
                    description=details,
                    balance_pence=parse_pence(cell(row, columns, "balance")),
                    kind=classify(details),
                )
            )
        else:
            cost = parse_pence(cell(row, columns, "cost"))
            if cost is None:
                statement.warnings.append(f"row {line_no}: trade with no cost, skipped")
                continue
            details = normalise_text(cell(row, columns, "details"))
            # Store sells as negative so summing cost gives net money put into funds
            cost = -abs(cost) if classify(details) == "sell" else abs(cost)
            statement.trades.append(
                InvestmentTrade(
                    traded_on=date,
                    fund=normalise_text(cell(row, columns, "fund")),
                    details=details,
                    quantity=Decimal(str(cell(row, columns, "quantity") or 0).replace(",", "")),
                    price=Decimal(str(cell(row, columns, "price") or 0).replace(",", "").replace("£", "")),
                    cost_pence=cost,
                )
            )

    if not statement.transactions and not statement.trades:
        raise ValueError(f"{path.name}: no Vanguard cash or investment transactions found")
    return statement
